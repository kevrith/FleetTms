"""M-Pesa statement import (masterplan 5.10). The owner uploads the business's M-Pesa statement (CSV or Excel). Money that
left the account is checked against what drivers claimed: a fuel entry, an expense or a float transfer with the same M-Pesa
code. A claim for more than the statement shows, or a claim with a code that is not on the statement at all, is flagged.
Money that came in from clients is matched to invoices (this also recovers a payment whose Safaricom notice was missed).

Payments to a lessor, loan repayments and salary advances are matched the same way, by the M-Pesa code recorded with them."""

import csv
import io
import re
import zipfile
from datetime import UTC, datetime, timedelta
from typing import Any

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models import (
    Expense,
    FinanceInstalment,
    FloatTransfer,
    FuelEntry,
    InvoicePayment,
    LeaseEntry,
    MpesaStatementImport,
    MpesaStatementLine,
    MpesaTransaction,
    SalaryAdvance,
)
from app.payments import record_payment
from app.reminders import NAIROBI

MAX_BYTES = 5 * 1024 * 1024
MAX_ROWS = 20000
CODE = re.compile(r"^[A-Z0-9]{8,12}$")
COLUMNS = {
    "receipt": ("receipt no.", "receipt no", "receipt number", "receipt", "transaction id", "transaction code"),
    "completed": ("completion time", "completed time", "completion date", "time", "date"),
    "details": ("details", "description", "transaction details"),
    "status": ("transaction status", "status"),
    "paid_in": ("paid in", "credit", "money in", "amount in"),
    "withdrawn": ("withdrawn", "paid out", "debit", "money out", "amount out"),
    "account": ("a/c no.", "a/c no", "account no", "account number", "bill ref number", "bill reference"),
}
DATE_FORMATS = ("%Y-%m-%d %H:%M:%S", "%d-%m-%Y %H:%M:%S", "%d/%m/%Y %H:%M:%S", "%d/%m/%Y %H:%M", "%Y-%m-%d %H:%M", "%d-%m-%Y %H:%M", "%Y/%m/%d %H:%M:%S")
FLAGS = ("statement_amount_differs", "not_on_statement")


class StatementError(Exception):
    """The file cannot be read as a statement. The text is safe to show."""


def _cells(content: bytes, filename: str | None) -> list[list[Any]]:
    if content[:2] == b"PK" or (filename or "").lower().endswith((".xlsx", ".xlsm")):
        try:
            with zipfile.ZipFile(io.BytesIO(content)):
                pass
            sheet = load_workbook(io.BytesIO(content), read_only=True, data_only=True).active
            return [list(r) for r in sheet.iter_rows(values_only=True)]
        except (zipfile.BadZipFile, KeyError, ValueError, OSError):
            raise StatementError("That Excel file could not be read. Export the statement again as CSV or Excel.") from None
    for encoding in ("utf-8-sig", "latin-1"):
        try:
            text = content.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    return [row for row in csv.reader(io.StringIO(text))]


def _find_header(rows: list[list[Any]]) -> tuple[int, dict[str, int]]:
    for n, row in enumerate(rows[:40]):
        names = [str(c or "").strip().lower() for c in row]
        found = {key: names.index(alias) for key, aliases in COLUMNS.items() for alias in aliases if alias in names}
        if "receipt" in found and "completed" in found and ("paid_in" in found or "withdrawn" in found):
            return n, found
    raise StatementError("This does not look like an M-Pesa statement. It needs columns for the receipt number, the completion time and the amount paid in or withdrawn.")


def _cents(raw: Any) -> int:
    if raw in (None, ""):
        return 0
    try:
        return abs(round(float(str(raw).replace(",", "").strip()) * 100))
    except ValueError:
        return 0


def _when(raw: Any) -> datetime | None:
    if isinstance(raw, datetime):
        return raw.replace(tzinfo=NAIROBI).astimezone(UTC) if raw.tzinfo is None else raw.astimezone(UTC)
    text = str(raw or "").strip()
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(text, fmt).replace(tzinfo=NAIROBI).astimezone(UTC)
        except ValueError:
            continue
    return None


def parse(content: bytes, filename: str | None = None) -> tuple[list[dict], int]:
    """(lines, number of rows skipped). Failed transactions, totals and anything without a receipt code are skipped."""
    if len(content) > MAX_BYTES:
        raise StatementError("Keep the file under 5 MB.")
    rows = _cells(content, filename)
    start, col = _find_header(rows)
    lines: list[dict] = []
    skipped = 0
    for row in rows[start + 1 :]:
        get = lambda key, row=row: row[col[key]] if key in col and col[key] < len(row) else None
        receipt = str(get("receipt") or "").strip().upper()
        when = _when(get("completed"))
        status = str(get("status") or "completed").strip().lower()
        if not CODE.match(receipt) or when is None or status not in ("completed", ""):
            skipped += bool(any(str(c or "").strip() for c in row))
            continue
        paid_in, withdrawn = _cents(get("paid_in")), _cents(get("withdrawn"))
        if paid_in == 0 and withdrawn == 0:
            skipped += 1
            continue
        lines.append({"receipt": receipt, "completed_at": when, "details": str(get("details") or "").strip()[:255] or None, "paid_in_cents": paid_in, "withdrawn_cents": withdrawn, "account": str(get("account") or "").strip() or None})
    if len(lines) > MAX_ROWS:
        raise StatementError("That statement has too many lines. Import it a month at a time.")
    if not lines:
        raise StatementError("No completed M-Pesa transactions were found in that file.")
    return lines, skipped


def _without(flags: list[str]) -> list[str]:
    return [f for f in flags if f not in FLAGS]


async def _match_out(db: AsyncSession, line: MpesaStatementLine) -> None:
    """Looks for what a line that left the account was for."""
    code = line.receipt
    for kind, model, when in (("fuel", FuelEntry, "captured_at"), ("expense", Expense, "spent_at"), ("float", FloatTransfer, "sent_at")):
        row = (await db.execute(select(model).where(model.mpesa_code == code))).scalar_one_or_none()
        if row is None:
            continue
        line.match_kind, line.match_id = kind, row.id
        if row.amount_cents == line.withdrawn_cents:
            line.state, line.note = "matched", None
            if hasattr(row, "flags"):
                row.flags = _without(row.flags)
        else:
            line.state = "amount_differs"
            line.note = f"The statement shows KES {line.withdrawn_cents / 100:,.2f} but KES {row.amount_cents / 100:,.2f} was claimed."
            if hasattr(row, "flags"):
                row.flags = [*_without(row.flags), "statement_amount_differs"]
        return
    lease = (await db.execute(select(LeaseEntry).where(LeaseEntry.mpesa_code == code))).scalar_one_or_none()
    loan = (await db.execute(select(FinanceInstalment).where(FinanceInstalment.mpesa_code == code))).scalar_one_or_none() if lease is None else None
    advance = (await db.execute(select(SalaryAdvance).where(SalaryAdvance.mpesa_code == code))).scalar_one_or_none() if lease is None and loan is None else None
    found = (("lease", lease.id, -lease.amount_cents) if lease else None) or (("finance", loan.id, loan.paid_cents) if loan else None) or (("advance", advance.id, advance.amount_cents) if advance else None)
    if found:
        line.match_kind, line.match_id = found[0], found[1]
        if found[2] == line.withdrawn_cents:
            line.state, line.note = "matched", None
        else:
            line.state, line.note = "amount_differs", f"The statement shows KES {line.withdrawn_cents / 100:,.2f} but KES {found[2] / 100:,.2f} was recorded."
        return
    line.match_kind = line.match_id = None
    line.state, line.note = "unmatched", "Money left the account and nothing in FleetTms claims this code."


async def _match_in(db: AsyncSession, line: MpesaStatementLine, account: str | None) -> bool:
    """Looks for the client payment a line that came in belongs to. Returns True if it created one."""
    txn = (await db.execute(select(MpesaTransaction).where(MpesaTransaction.trans_id == line.receipt))).scalar_one_or_none()
    created = False
    if txn is None:
        by_hand = (await db.execute(select(InvoicePayment).where(InvoicePayment.reference == line.receipt))).scalars().first()
        if by_hand is not None:
            line.match_kind, line.match_id, line.state, line.note = "client_payment", by_hand.id, "matched", None
            return False
        txn, created = await record_payment(
            db, trans_id=line.receipt, amount_cents=line.paid_in_cents, bill_ref=account, payer_name=(line.details or "")[:160], payer_phone=None,
            shortcode=None, paid_at=line.completed_at, source="statement",
        )  # fmt: skip
    line.match_kind, line.match_id = "client_payment", txn.id
    if txn.status == "matched":
        line.state, line.note = "matched", None
    else:
        line.state, line.note = "unmatched", "A client paid in. Match it to an invoice in Payments."
    return created


async def import_statement(db: AsyncSession, lines: list[dict], skipped: int, *, filename: str | None, user_id) -> MpesaStatementImport:
    imported = MpesaStatementImport(filename=(filename or "")[:200] or None, rows=len(lines), created_by_user_id=user_id)
    db.add(imported)
    await db.flush()
    known = {r.receipt: r for r in (await db.execute(select(MpesaStatementLine).where(MpesaStatementLine.receipt.in_([x["receipt"] for x in lines])))).scalars()}
    new = recovered = 0
    touched: list[MpesaStatementLine] = []
    for item in lines:
        line = known.get(item["receipt"])
        if line is None:
            line = MpesaStatementLine(
                import_id=imported.id, receipt=item["receipt"], completed_at=item["completed_at"], details=item["details"], paid_in_cents=item["paid_in_cents"],
                withdrawn_cents=item["withdrawn_cents"],
            )  # fmt: skip
            db.add(line)
            await db.flush()
            new += 1
        elif line.state in ("matched", "ignored"):
            continue  # already dealt with
        touched.append(line)
        if line.withdrawn_cents > 0:
            await _match_out(db, line)
        else:
            recovered += await _match_in(db, line, item["account"])
    await db.flush()

    # The other way round: claims in this period whose M-Pesa code is not on any statement we hold.
    first, last = min(x["completed_at"] for x in lines), max(x["completed_at"] for x in lines)
    window_start, window_end = first + timedelta(hours=6), last - timedelta(hours=6)
    on_statement = {r for (r,) in (await db.execute(select(MpesaStatementLine.receipt))).all()}
    missing: list[dict] = []
    for kind, model, when in (("fuel", FuelEntry, FuelEntry.captured_at), ("expense", Expense, Expense.spent_at), ("float", FloatTransfer, FloatTransfer.sent_at)):
        for row in (await db.execute(select(model).where(model.mpesa_code.is_not(None), when >= window_start, when <= window_end))).scalars():
            if row.mpesa_code in on_statement:
                continue
            if hasattr(row, "flags") and "not_on_statement" not in row.flags:
                row.flags = [*row.flags, "not_on_statement"]
            missing.append({"kind": kind, "id": str(row.id), "mpesa_code": row.mpesa_code, "amount_cents": row.amount_cents})
    states = [ln.state for ln in touched]
    imported.new_rows, imported.period_start, imported.period_end = new, first, last
    imported.summary = {
        "skipped": skipped, "already_known": len(lines) - new, "matched": states.count("matched"), "amount_differs": states.count("amount_differs"),
        "unmatched": states.count("unmatched"), "payments_recovered": recovered, "not_on_statement": missing,
    }  # fmt: skip
    return imported
