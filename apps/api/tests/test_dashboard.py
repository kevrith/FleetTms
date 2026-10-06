from datetime import timedelta

from app.reminders import nairobi_today
from tests.fleet import add_vehicle
from tests.helpers import (
    PASSWORD,
    bearer,
    driver_session,
    enable_2fa,
    login,
    owner_session,
    staff_session,
)
from tests.shots import begin_trip, fleet, inspect, send_float


async def dash(client, tokens):
    res = await client.get("/dashboard", headers=bearer(tokens))
    assert res.status_code == 200, res.text
    return res.json()


def kinds(d):
    return [a["kind"] for a in d["alerts"]]


async def finish_trip(client, f, trip, end=125580):
    from tests.shots import photo_id

    pid = await photo_id(client, f.driver, "odometer")
    res = await client.post(f"/trips/{trip['id']}/end", headers=bearer(f.driver), json={"photo_id": pid, "value": end})
    assert res.status_code == 200, res.text


async def test_todays_numbers_add_up(client):
    f = await fleet(client)
    await send_float(client, f, 300000)
    trip = await begin_trip(client, f)
    await client.post("/expenses", headers=bearer(f.driver), json={"category": "toll", "amount_cents": 50000})
    await client.post("/fuel", headers=bearer(f.driver), json={"vehicle_id": f.vehicle["id"], "litres": "120", "price_per_litre_cents": 18500, "amount_cents": 2220000})
    d = await dash(client, f.owner)
    n = d["numbers"]
    assert n["trips_active"] == 1 and n["trips_completed_today"] == 0
    assert n["expenses_today_cents"] == 50000 and n["floats_sent_today_cents"] == 300000
    assert n["fuel_today"] == {"litres": 120.0, "amount_cents": 2220000}
    await finish_trip(client, f, trip)
    n = (await dash(client, f.owner))["numbers"]
    assert n["trips_active"] == 0 and n["trips_completed_today"] == 1 and n["distance_today_km"] == 480
    assert n["income_today_cents"] == 0 and n["money_owed_cents"] == 0  # nothing invoiced yet (billing is Sprint 8 and 9)
    assert n["mode"] == "standard" and n["day"] == nairobi_today().isoformat()


async def test_attention_list_shows_what_needs_doing_with_red_before_amber(client):
    f = await fleet(client)
    await send_float(client, f, 300000)
    await inspect(client, f.driver, f.vehicle["id"], {"Brakes": "Pedal goes to the floor"})  # blocked + urgent work order
    await client.post("/documents", headers=bearer(f.owner), json={"doc_type": "insurance", "vehicle_id": f.vehicle["id"], "expires_on": (nairobi_today() - timedelta(days=2)).isoformat()})
    await client.post(f"/vehicles/{f.vehicle['id']}/services", headers=bearer(f.owner), json={"name": "Oil", "every_km": 10000, "last_done_km": 110000})  # 15,000 km since: overdue
    await client.put("/spend-limits", headers=bearer(f.owner), json=[{"limit_cents": 10000}])
    await client.post("/expenses", headers=bearer(f.driver), json={"category": "toll", "amount_cents": 50000})  # over the limit
    await client.post("/fuel", headers=bearer(f.driver), json={"vehicle_id": f.vehicle["id"], "litres": "120", "price_per_litre_cents": 18500, "amount_cents": 3000000})
    await client.post("/devices/integrity", headers=bearer(f.driver), json={"device_id": "pixel-1", "vehicle_id": f.vehicle["id"], "mock_location": True})
    await client.post("/me/reconciliation/submit", headers=bearer(f.driver), json={})
    d = await dash(client, f.owner)
    got = set(kinds(d))
    assert {"inspection_blocked", "work_order_urgent", "document_expired", "service_overdue", "expense_waiting", "fuel_flagged", "low_trust", "reconciliation_waiting"} <= got
    severities = [a["severity"] for a in d["alerts"]]
    assert severities == sorted(severities, key=lambda s: {"red": 0, "amber": 1}[s])  # red first
    assert d["open_defects"] == 1
    blocked = next(a for a in d["alerts"] if a["kind"] == "inspection_blocked")
    assert blocked["severity"] == "red" and "KCA 123A" in blocked["title"] and blocked["link"] == f"/vehicles/{f.vehicle['id']}"


async def test_a_quiet_business_has_no_alerts(client):
    f = await fleet(client)
    d = await dash(client, f.owner)
    assert d["alerts"] == [] and d["open_defects"] == 0


async def test_an_owner_sees_who_is_in_trouble_first_with_a_way_to_call(client):
    f = await fleet(client)
    assert (await dash(client, f.owner))["emergencies"] == []
    await client.post("/sos", headers=bearer(f.driver), json={"lat": -1.29, "lng": 36.82})
    await client.post("/incidents", headers=bearer(f.driver), json={"type": "breakdown", "description": "Engine overheated", "lat": -1.3, "lng": 36.8})
    await client.post("/incidents", headers=bearer(f.driver), json={"type": "traffic_fine", "description": "Speeding"})
    d = await dash(client, f.owner)
    assert sorted(e["kind"] for e in d["emergencies"]) == ["breakdown", "sos"]
    sos = next(e for e in d["emergencies"] if e["kind"] == "sos")
    assert sos["registration"] == "KCA 123A" and sos["answered"] is False and "phone" in sos and sos["lat"] == -1.29
    assert next(e for e in d["emergencies"] if e["kind"] == "breakdown")["registration"] == "KCA 123A"
    flagged = [a["kind"] for a in d["alerts"] if a.get("emergency")]
    assert sorted(flagged) == ["incident_open", "sos_active"]  # the fine is an ordinary alert, not an emergency
    accountant, _ = await staff_session(client, f.owner, "accountant", "acc2@example.com")
    assert (await dash(client, accountant))["emergencies"] == []  # the accountant neither answers an SOS nor handles incidents


async def test_an_unanswered_emergency_comes_before_an_answered_one_and_an_sos_before_a_breakdown(client):
    f = await fleet(client)
    sos = (await client.post("/sos", headers=bearer(f.driver), json={"lat": -1.29, "lng": 36.82})).json()
    await client.post("/incidents", headers=bearer(f.driver), json={"type": "breakdown", "description": "Engine overheated"})
    order = [(e["kind"], e["answered"]) for e in (await dash(client, f.owner))["emergencies"]]
    assert order == [("sos", False), ("breakdown", False)]
    await client.post(f"/sos/{sos['id']}/acknowledge", headers=bearer(f.owner), json={})
    order = [(e["kind"], e["answered"]) for e in (await dash(client, f.owner))["emergencies"]]
    assert order == [("breakdown", False), ("sos", True)]  # the one nobody has answered goes first


async def test_documents_and_tyres_are_worded_the_way_an_owner_would_say_them(client):
    f = await fleet(client)
    today = nairobi_today()
    for doc, days in (("insurance", 4), ("inspection", -2)):
        await client.post("/documents", headers=bearer(f.owner), json={"doc_type": doc, "vehicle_id": f.vehicle["id"], "expires_on": (today + timedelta(days=days)).isoformat()})
    titles = [a["title"] for a in (await dash(client, f.owner))["alerts"]]
    assert "KCA 123A: insurance expires in 4 days" in titles and "KCA 123A: inspection expired 2 days ago" in titles


def test_days_are_said_in_plain_words():
    from datetime import date

    from app.routers.dashboard import days_phrase

    today = date(2026, 10, 5)
    said = [days_phrase(date(2026, 10, d), today) for d in (5, 6, 9, 4, 1)]
    assert said == ["today", "tomorrow", "in 4 days", "yesterday", "4 days ago"]


def test_an_old_work_order_title_loses_its_underscores():
    from app.wording import plain_title

    assert plain_title("Tyre (steer_left): do not drive") == "Tyre (steer left): do not drive"
    assert plain_title("Tyre (rear_axle_outer_left): driver request") == "Tyre (rear axle outer left): driver request"
    assert plain_title("Brakes: do not drive") == "Brakes: do not drive"
    assert plain_title("Fix the some_thing sensor") == "Fix the some_thing sensor"


async def test_the_fleet_strip_counts_every_active_lorry_once(client):
    f = await fleet(client)
    await add_vehicle(client, f.owner, "KCB 222B")
    d = await dash(client, f.owner)
    counts = d["fleet"]
    assert counts["total"] == 2
    assert sum(counts[k] for k in ("moving", "idle", "offline", "parked", "unknown", "in_workshop")) == 2
    await inspect(client, f.driver, f.vehicle["id"], {"Body damage": "Dent on the door"})
    wo = (await client.get("/work-orders", headers=bearer(f.owner))).json()[0]
    await client.put(f"/work-orders/{wo['id']}", headers=bearer(f.owner), json={"status": "in_progress", "assignee_kind": "garage", "assignee_name": "Kamau Motors"})
    assert (await dash(client, f.owner))["fleet"]["in_workshop"] == 1


async def test_each_person_sees_only_their_part(client):
    f = await fleet(client)
    await client.put("/spend-limits", headers=bearer(f.owner), json=[{"limit_cents": 10000}])
    await client.post("/expenses", headers=bearer(f.driver), json={"category": "toll", "amount_cents": 50000})
    await inspect(client, f.driver, f.vehicle["id"], {"Brakes": "Soft"})
    manager, _ = await staff_session(client, f.owner, "manager", "m@example.com")
    assert "expense_waiting" not in kinds(await dash(client, manager))  # only the owner can approve over-limit spending
    assert "inspection_blocked" in kinds(await dash(client, manager))
    accountant, _ = await staff_session(client, f.owner, "accountant", "acc@example.com")
    acc = await dash(client, accountant)
    assert "inspection_blocked" not in kinds(acc) and "expenses_today_cents" in acc["numbers"] and "fuel_today" not in acc["numbers"]
    assert (await client.get("/dashboard", headers=bearer(f.driver))).status_code == 403
    workshop, _ = await staff_session(client, f.owner, "workshop", "w@example.com", with_2fa=False)
    assert (await client.get("/dashboard", headers=bearer(workshop))).status_code == 403
    other, _ = await owner_session(client, "Bravo", "b@example.com")
    assert (await dash(client, other))["alerts"] == []


async def test_a_supervisor_sees_alerts_for_their_vehicles_only(client):
    f = await fleet(client)
    other_vehicle = await add_vehicle(client, f.owner, "KCB 222B")
    for v in (f.vehicle, other_vehicle):
        await client.post("/documents", headers=bearer(f.owner), json={"doc_type": "insurance", "vehicle_id": v["id"], "expires_on": (nairobi_today() - timedelta(days=1)).isoformat()})
    invite = await client.post("/users", headers=bearer(f.owner), json={"name": "Sup", "email": "sup@example.com", "roles": ["supervisor"], "vehicle_scope": [f.vehicle["id"]]})
    await client.post("/auth/accept-invite", json={"token": invite.json()["invite_token"], "password": PASSWORD})
    sup = (await login(client, "sup@example.com")).json()
    await enable_2fa(client, sup)
    titles = [a["title"] for a in (await dash(client, sup))["alerts"]]
    assert len(titles) == 1 and "KCA 123A" in titles[0]


async def test_owner_driver_mode_skips_alerts_about_their_own_entries(client):
    owner, _ = await owner_session(client)
    vehicle = (await client.post("/vehicles", headers=bearer(owner), json={"registration": "KCA 777A", "odometer_km": 1000})).json()
    me = next(s for s in (await client.get("/staff", headers=bearer(owner))).json())
    await client.put(f"/users/{me['membership_id']}/roles", headers=bearer(owner), json={"roles": ["owner", "driver"]})
    await client.post(f"/vehicles/{vehicle['id']}/crew", headers=bearer(owner), json={"membership_id": me["membership_id"], "role": "driver"})
    await client.post("/floats", headers=bearer(owner), json={"driver_membership_id": me["membership_id"], "amount_cents": 100000})
    await client.post("/me/reconciliation/submit", headers=bearer(owner), json={})
    d = await dash(client, owner)
    assert d["numbers"]["mode"] == "owner_driver"
    assert "reconciliation_waiting" not in kinds(d)  # nobody else needs to approve the owner's own day
    own_expense = await client.post("/expenses", headers=bearer(owner), json={"category": "toll", "amount_cents": 9_000_000})
    assert own_expense.json()["status"] == "recorded"  # and the owner is never made to approve their own spending


# ---- reports ---------------------------------------------------------------------------------------


async def test_the_daily_report_totals_trips_fuel_and_expenses_by_category_and_vehicle(client):
    f = await fleet(client)
    await send_float(client, f, 300000)
    trip = await begin_trip(client, f)
    await client.post("/expenses", headers=bearer(f.driver), json={"category": "toll", "amount_cents": 50000})
    await client.post("/expenses", headers=bearer(f.driver), json={"category": "food", "amount_cents": 30000})
    await client.post("/expenses", headers=bearer(f.driver), json={"category": "toll", "amount_cents": 20000})
    await client.post("/fuel", headers=bearer(f.driver), json={"vehicle_id": f.vehicle["id"], "litres": "100", "price_per_litre_cents": 18500, "amount_cents": 1850000})
    await finish_trip(client, f, trip)
    today = nairobi_today().isoformat()
    res = await client.get(f"/reports/summary?from={today}&to={today}", headers=bearer(f.owner))
    assert res.status_code == 200, res.text
    r = res.json()
    assert r["totals"] == {"trips_completed": 1, "distance_km": 480, "fuel_litres": 100.0, "fuel_cents": 1850000, "expenses_cents": 100000, "floats_sent_cents": 300000}
    assert r["expenses_by_category"] == [{"category": "toll", "cents": 70000}, {"category": "food", "cents": 30000}]
    row = r["by_vehicle"][0]
    assert row["registration"] == "KCA 123A" and row["trips"] == 1 and row["distance_km"] == 480 and row["km_per_litre"] == 4.8
    assert [d["day"] for d in r["by_day"]] == [today] and r["by_day"][0]["expenses_cents"] == 100000


async def test_a_weekly_report_lists_every_day_even_quiet_ones(client):
    f = await fleet(client)
    today = nairobi_today()
    start = today - timedelta(days=6)
    r = (await client.get(f"/reports/summary?from={start.isoformat()}&to={today.isoformat()}", headers=bearer(f.owner))).json()
    assert len(r["by_day"]) == 7 and r["totals"]["trips_completed"] == 0 and r["by_vehicle"] == []


async def test_report_ranges_and_permissions(client):
    f = await fleet(client)
    today = nairobi_today()
    bad = await client.get(f"/reports/summary?from={today.isoformat()}&to={(today - timedelta(days=1)).isoformat()}", headers=bearer(f.owner))
    assert bad.json()["detail"]["code"] == "bad_range"
    long = await client.get(f"/reports/summary?from={(today - timedelta(days=400)).isoformat()}&to={today.isoformat()}", headers=bearer(f.owner))
    assert long.json()["detail"]["code"] == "range_too_long"
    q = f"?from={today.isoformat()}&to={today.isoformat()}"
    manager, _ = await staff_session(client, f.owner, "manager", "m@example.com")
    accountant, _ = await staff_session(client, f.owner, "accountant", "acc@example.com")
    assert (await client.get("/reports/summary" + q, headers=bearer(manager))).status_code == 200
    assert (await client.get("/reports/summary" + q, headers=bearer(accountant))).status_code == 200
    assert (await client.get("/reports/summary" + q, headers=bearer(f.driver))).status_code == 403
    stranger = await driver_session(client, f.owner, "0733345678")
    assert (await client.get("/reports/summary" + q, headers=bearer(stranger))).status_code == 403
    other, _ = await owner_session(client, "Bravo", "b@example.com")
    assert (await client.get("/reports/summary" + q, headers=bearer(other))).json()["totals"]["expenses_cents"] == 0


async def test_a_report_can_be_downloaded_as_excel_or_pdf(client):
    from io import BytesIO

    from openpyxl import load_workbook

    f = await fleet(client)
    await client.post("/expenses", headers=bearer(f.driver), json={"category": "toll", "amount_cents": 50000})
    today = nairobi_today().isoformat()
    q = f"?from={today}&to={today}"
    xlsx = await client.get("/reports/export" + q, headers=bearer(f.owner))
    assert xlsx.status_code == 200 and "spreadsheetml" in xlsx.headers["content-type"]
    book = load_workbook(BytesIO(xlsx.content))
    assert book.sheetnames == ["Totals", "Expenses by type", "By vehicle", "By day"]
    assert ["Toll", 500] in [list(r) for r in book["Expenses by type"].iter_rows(min_row=2, values_only=True)]
    pdf = await client.get("/reports/export" + q + "&format=pdf", headers=bearer(f.owner))
    assert pdf.status_code == 200 and pdf.content.startswith(b"%PDF")
    assert (await client.get("/reports/export" + q, headers=bearer(f.driver))).status_code == 403


async def test_sprint_6_alerts_reach_the_right_people_with_sos_first(client):
    from tests.test_parts import issue, part, receive, work_order
    from tests.test_tyres import fit, tyre

    f = await fleet(client)
    t = await tyre(client, f.owner, "A1111")
    await fit(client, f.owner, t, f.vehicle, "steer_left")
    from tests.test_tyres import inspect_with

    await inspect_with(client, f, [{"position": "steer_left", "serial": "ZZ9999"}])
    p = await receive(client, f.owner, await part(client, f.owner, reorder_level=5), 6, 100000)
    wo = await work_order(client, f)
    await issue(client, f, wo, p, 2)
    await client.post(f"/work-orders/{wo['id']}/complete", headers=bearer(f.owner), json={})
    await client.post("/incidents", headers=bearer(f.driver), json={"type": "accident", "description": "Hit a pothole"})
    await client.post("/incidents", headers=bearer(f.driver), json={"type": "police_stop"})
    await client.post("/sos", headers=bearer(f.driver), json={"lat": -1.29, "lng": 36.82})
    owner_view = await dash(client, f.owner)
    assert owner_view["alerts"][0]["kind"] == "sos_active" and owner_view["alerts"][0]["severity"] == "red"
    got = kinds(owner_view)
    for kind in ("tyre_swap", "stock_low", "parts_unfitted", "incident_open"):
        assert kind in got, got
    severities = {a["kind"]: a["severity"] for a in owner_view["alerts"]}
    assert severities["tyre_swap"] == "red" and severities["stock_low"] == "amber"
