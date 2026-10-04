"""Excel import of clients, suppliers, and what clients still owed when the business started."""

from datetime import timedelta

from app.reminders import nairobi_today
from app.routers.imports import BALANCE_COLUMNS, CLIENT_COLUMNS, SUPPLIER_COLUMNS
from tests.helpers import bearer, driver_session, owner_session
from tests.test_import import upload, workbook

TODAY = nairobi_today()


def client_row(name, **extra):
    return {"name": name, "contact_name": "Grace Mwangi", "phone": "0712345678", "billing_method": "per_tonne", "rate_kes": 3000, "payment_terms_days": 45, "vat_pct": 16, **extra}


def balance_row(client, number, amount, **extra):
    return {"client": client, "invoice_number": number, "invoice_date": (TODAY - timedelta(days=60)).isoformat(), "due_date": (TODAY - timedelta(days=30)).isoformat(), "amount_kes": amount, "note": "Cement to Nairobi", **extra}


async def import_ok(client, owner, kind, columns, rows):
    res = await upload(client, owner, kind, workbook(columns, rows))
    assert res.status_code == 200, res.text
    assert res.json()["errors"] == [] and res.json()["imported"] == len(rows), res.text
    return res.json()


# ---- clients -------------------------------------------------------------------------------------------------------------


async def test_clients_import_with_their_rates_terms_and_vat_in_the_right_units(client):
    owner, _ = await owner_session(client)
    await import_ok(client, owner, "clients", CLIENT_COLUMNS, [client_row("Mwangi Cement Ltd", kra_pin="P051234567Z", email="Accounts@Mwangi.example.com"), client_row("Rift Freight", rate_kes="1250.50", billing_method="per_trip")])
    clients = {c["name"]: c for c in (await client.get("/clients", headers=bearer(owner))).json()}
    m = clients["Mwangi Cement Ltd"]
    assert m["rate_cents"] == 300_000 and m["billing_method"] == "per_tonne" and m["payment_terms_days"] == 45 and float(m["vat_pct"]) == 16 and m["phone"] == "+254712345678"
    assert clients["Rift Freight"]["rate_cents"] == 125_050 and clients["Rift Freight"]["billing_method"] == "per_trip"
    audit = [a for a in (await client.get("/audit", params={"action": "client.added"}, headers=bearer(owner))).json()]
    assert len(audit) == 2 and all(a["note"] == "Created by Excel import" for a in audit)


async def test_a_clients_file_with_problems_imports_nothing_and_lists_every_problem_by_row(client):
    owner, _ = await owner_session(client)
    await import_ok(client, owner, "clients", CLIENT_COLUMNS, [client_row("Existing Client")])
    rows = [
        client_row("Good Client"),
        client_row("existing client"),  # already there, whatever the capitals
        client_row("Good Client"),  # twice in the file
        client_row("Bad Phone", phone="12"),
        client_row("Bad Method", billing_method="by the moon"),
        client_row("Fractions", rate_kes="10.005"),
        client_row("   "),
    ]
    res = await upload(client, owner, "clients", workbook(CLIENT_COLUMNS, rows))
    body = res.json()
    assert body["imported"] == 0
    by_row = {e["row"]: e for e in body["errors"]}
    assert set(by_row) == {3, 4, 5, 6, 7, 8}  # every bad row, numbered as in the sheet; row 2 was fine
    assert by_row[7]["column"] == "rate_kes" and "two decimals" in by_row[7]["message"]
    assert by_row[3]["column"] == "name" and by_row[4]["column"] == "name"
    assert [c["name"] for c in (await client.get("/clients", headers=bearer(owner))).json()] == ["Existing Client"]  # all or nothing


async def test_a_dry_run_checks_the_file_and_changes_nothing(client):
    owner, _ = await owner_session(client)
    res = await upload(client, owner, "clients", workbook(CLIENT_COLUMNS, [client_row("Dry Run Ltd")]), dry_run=True)
    assert res.json()["errors"] == [] and res.json()["imported"] == 0 and res.json()["rows"] == 1
    assert (await client.get("/clients", headers=bearer(owner))).json() == []


# ---- suppliers ---------------------------------------------------------------------------------------------------------------


async def test_suppliers_import_with_clean_contacts_and_refuse_duplicates_and_bad_emails(client):
    owner, _ = await owner_session(client)
    await import_ok(client, owner, "suppliers", SUPPLIER_COLUMNS, [{"name": "Kiambu Spares", "phone": "0722345678", "category": "spares", "email": "sales@kiambu.example.com"}, {"name": "Nakuru Tyres"}])
    suppliers = {s["name"]: s for s in (await client.get("/suppliers", headers=bearer(owner))).json()}
    assert suppliers["Kiambu Spares"]["phone"] == "+254722345678" and suppliers["Kiambu Spares"]["category"] == "spares" and "Nakuru Tyres" in suppliers
    res = await upload(client, owner, "suppliers", workbook(SUPPLIER_COLUMNS, [{"name": "kiambu spares"}, {"name": "Fresh One", "email": "not an email"}, {"name": "Fresh Two"}, {"name": "fresh two"}]))
    assert res.json()["imported"] == 0
    assert {e["row"] for e in res.json()["errors"]} == {2, 3, 5}
    assert len((await client.get("/suppliers", headers=bearer(owner))).json()) == 2


# ---- what clients owed when the business started -----------------------------------------------------------------------


async def test_opening_balances_become_unpaid_invoices_that_show_up_as_debt_with_their_age(client):
    owner, _ = await owner_session(client)
    await import_ok(client, owner, "clients", CLIENT_COLUMNS, [client_row("Mwangi Cement Ltd"), client_row("Rift Freight")])
    rows = [balance_row("Mwangi Cement Ltd", "INV-0412", 245_000), balance_row("mwangi cement ltd", "INV-0420", "50000.50"), balance_row("Rift Freight", "9981", 80_000, due_date=(TODAY + timedelta(days=10)).isoformat())]
    await import_ok(client, owner, "balances", BALANCE_COLUMNS, rows)

    invoices = (await client.get("/invoices", headers=bearer(owner))).json()
    assert sorted(i["number"] for i in invoices) == ["OB-9981", "OB-INV-0412", "OB-INV-0420"]
    first = next(i for i in invoices if i["number"] == "OB-INV-0412")
    assert first["kind"] == "opening" and first["total_cents"] == 24_500_000 and first["status"] == "issued"

    debtors = {d["name"]: d for d in (await client.get("/debtors", headers=bearer(owner))).json()["clients"]}
    mwangi, rift = debtors["Mwangi Cement Ltd"], debtors["Rift Freight"]
    assert mwangi["balance_cents"] == 24_500_000 + 5_000_050 and mwangi["overdue_cents"] == mwangi["balance_cents"] and mwangi["invoices"] == 2
    assert mwangi["buckets"]["1_30"] == mwangi["balance_cents"]  # due 30 days ago: the first late bucket
    assert rift["balance_cents"] == 8_000_000 and rift["overdue_cents"] == 0  # not yet due

    detail = (await client.get(f"/invoices/{first['id']}", headers=bearer(owner))).json()
    assert detail["lines"][0]["description"].startswith("Balance brought forward: invoice INV-0412")
    assert (await client.get(f"/invoices/{first['id']}/pdf", headers=bearer(owner))).content.startswith(b"%PDF")  # it can be printed and sent like any invoice

    paid = await client.post(f"/invoices/{first['id']}/payments", headers=bearer(owner), json={"amount_cents": 24_500_000, "method": "bank", "reference": "RTGS12345"})
    assert paid.status_code == 201, paid.text  # a client paying the old invoice is recorded like any payment
    after = {d["name"]: d for d in (await client.get("/debtors", headers=bearer(owner))).json()["clients"]}
    assert after["Mwangi Cement Ltd"]["balance_cents"] == 5_000_050


async def test_an_opening_balance_is_not_income_and_is_not_sent_to_the_tax_authority(client, monkeypatch):
    from app import etims_service

    async def never(*args, **kwargs):
        raise AssertionError("an opening balance must not be queued for eTIMS")

    monkeypatch.setattr(etims_service, "queue_sale", never)
    owner, _ = await owner_session(client)
    await import_ok(client, owner, "clients", CLIENT_COLUMNS, [client_row("Mwangi Cement Ltd")])
    await import_ok(client, owner, "balances", BALANCE_COLUMNS, [balance_row("Mwangi Cement Ltd", "INV-1", 100_000, invoice_date=TODAY.isoformat(), due_date=(TODAY + timedelta(days=30)).isoformat())])
    numbers = (await client.get("/dashboard", headers=bearer(owner))).json()["numbers"]
    assert numbers["income_today_cents"] == 0  # dated today, but invoiced and counted on the old system
    assert numbers["money_owed_cents"] == 10_000_000


async def test_a_balances_file_with_problems_imports_nothing(client):
    owner, _ = await owner_session(client)
    await import_ok(client, owner, "clients", CLIENT_COLUMNS, [client_row("Mwangi Cement Ltd")])
    await import_ok(client, owner, "balances", BALANCE_COLUMNS, [balance_row("Mwangi Cement Ltd", "OLD-1", 1000)])
    rows = [
        balance_row("Nobody Ltd", "A-1", 1000),  # no such client
        balance_row("Mwangi Cement Ltd", "OLD-1", 1000),  # already imported
        balance_row("Mwangi Cement Ltd", "A-2", 1000, due_date=(TODAY - timedelta(days=90)).isoformat()),  # due before it was issued
        balance_row("Mwangi Cement Ltd", "A-3", 0),
        balance_row("Mwangi Cement Ltd", "A-4", "10.005"),
        balance_row("Mwangi Cement Ltd", "", 1000),
        balance_row("Mwangi Cement Ltd", "A-5" * 8, 1000),
        balance_row("Mwangi Cement Ltd", "A-6", 1000, invoice_date=None),
        balance_row("Mwangi Cement Ltd", "A-7", 1000),
        balance_row("Mwangi Cement Ltd", "A-7", 1000),  # twice in the file
    ]
    body = (await upload(client, owner, "balances", workbook(BALANCE_COLUMNS, rows))).json()
    assert body["imported"] == 0
    cols = {e["row"]: e["column"] for e in body["errors"]}
    assert cols[2] == "client" and cols[3] == "invoice_number" and cols[4] == "due_date" and cols[5] == "amount_kes" and cols[6] == "amount_kes"
    assert cols[7] == "invoice_number" and cols[8] == "invoice_number" and cols[9] == "invoice_date" and cols[11] == "invoice_number" and 10 not in cols
    assert len((await client.get("/invoices", headers=bearer(owner))).json()) == 1  # only the earlier good import


# ---- templates and permissions ----------------------------------------------------------------------------------------------


async def test_every_new_import_has_a_template_with_its_columns_and_an_example_row(client):
    from io import BytesIO

    from openpyxl import load_workbook

    owner, _ = await owner_session(client)
    for kind, columns in (("clients", CLIENT_COLUMNS), ("suppliers", SUPPLIER_COLUMNS), ("balances", BALANCE_COLUMNS)):
        res = await client.get(f"/imports/{kind}/template", headers=bearer(owner))
        assert res.status_code == 200
        rows = list(load_workbook(BytesIO(res.content)).active.iter_rows(values_only=True))
        assert list(rows[0]) == columns and len(rows) == 2


async def test_only_people_allowed_to_import_can_and_the_wrong_columns_are_refused(client):
    owner, _ = await owner_session(client)
    driver = await driver_session(client, owner, "0733333333")
    for kind in ("clients", "suppliers", "balances"):
        assert (await upload(client, driver, kind, b"x")).status_code == 403
        assert (await client.get(f"/imports/{kind}/template", headers=bearer(driver))).status_code == 403
    wrong = await upload(client, owner, "clients", workbook(["name", "colour"], [{"name": "X Ltd", "colour": "red"}]))
    assert wrong.status_code == 422 and wrong.json()["detail"]["code"] == "bad_columns"
