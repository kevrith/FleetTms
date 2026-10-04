import io
from datetime import date, timedelta

import pytest
from openpyxl import load_workbook

from app.reminders import nairobi_today
from app.report_delivery import fake_sender
from app.report_schedules import send_scheduled_reports
from tests.fraud_scenarios import drive, seeded
from tests.helpers import bearer, owner_session, staff_session
from tests.leasing import lease_setup, moment, seed_client, seed_trip
from tests.shots import fleet
from tests.test_predictions import history


@pytest.fixture(autouse=True)
def _empty_outboxes():
    for channel in ("email", "whatsapp"):
        fake_sender(channel).outbox.clear()
        fake_sender(channel).fail_with = None


def span(days=60):
    today = nairobi_today()
    return {"start": (today - timedelta(days=days)).isoformat(), "end": today.isoformat()}


async def run(client, who, key, **params):
    res = await client.get(f"/report-catalog/{key}", params=params or span(), headers=bearer(who))
    assert res.status_code == 200, res.text
    return res.json()


def section(report, title):
    return next(s for s in report["sections"] if s["title"].startswith(title))


# ---- who sees which report ----------------------------------------------------------------------------------------------------


async def test_each_role_sees_the_reports_it_may_run(client):
    f = await fleet(client)
    owner = (await client.get("/report-catalog", headers=bearer(f.owner))).json()
    assert [r["key"] for r in owner] == ["summary", "profit", "fuel_efficiency", "debtors", "alerts", "scorecards"] and all(r["available"] for r in owner)
    accountant, _ = await staff_session(client, f.owner, "accountant", "acc@example.com")
    assert [r["key"] for r in (await client.get("/report-catalog", headers=bearer(accountant))).json()] == ["summary", "profit", "fuel_efficiency", "debtors", "scorecards"]
    manager, _ = await staff_session(client, f.owner, "manager", "mgr@example.com")
    assert [r["key"] for r in (await client.get("/report-catalog", headers=bearer(manager))).json()] == ["summary", "fuel_efficiency", "alerts", "scorecards"]
    assert (await client.get("/report-catalog", headers=bearer(f.driver))).json() == []
    assert (await client.get("/report-catalog/profit", headers=bearer(manager))).status_code == 403
    assert (await client.get("/report-catalog/summary", headers=bearer(f.driver))).status_code == 403
    assert (await client.get("/report-catalog/nonsense", headers=bearer(f.owner))).status_code == 404
    assert (await client.get("/report-catalog")).status_code == 401


async def test_the_period_is_checked(client):
    f = await fleet(client)
    bad = await client.get("/report-catalog/summary", params={"start": "2026-10-10", "end": "2026-10-01"}, headers=bearer(f.owner))
    assert bad.status_code == 422 and bad.json()["detail"]["code"] == "bad_range"
    long = await client.get("/report-catalog/summary", params={"start": "2024-01-01", "end": "2026-10-01"}, headers=bearer(f.owner))
    assert long.json()["detail"]["code"] == "range_too_long"
    res = await client.get("/report-catalog/summary", headers=bearer(f.owner))
    assert res.status_code == 200
    default = res.json()
    assert (date.fromisoformat(default["to"]) - date.fromisoformat(default["from"])).days == 29  # the last 30 days


# ---- the reports -----------------------------------------------------------------------------------------------------------------


async def test_the_profit_report_agrees_with_the_profit_pages(client):
    owner, vehicle, _ = await lease_setup(client)
    c = await seed_client()
    for offset in (-2, -1):
        await seed_trip(vehicle["id"], revenue_cents=10_000_000, client_id=c, at=moment(offset, 15), km=400)
    start, end = (nairobi_today().replace(day=1) - timedelta(days=62)).isoformat(), nairobi_today().isoformat()
    report = await run(client, owner, "profit", start=start, end=end)
    business = {row[0]: row[1] for row in section(report, "The business")["rows"]}
    page = (await client.get("/profit", params={"from_month": start, "to_month": end}, headers=bearer(owner))).json()["business"]
    assert business["Revenue"] == page["revenue"] / 100 == 200_000.0
    assert business["Net profit across vehicles"] == page["net"] / 100 and business["Net profit after overheads"] == page["net_after_overheads"] / 100
    [row] = section(report, "By vehicle")["rows"]
    assert row[0] == "KDB 404D" and row[1] == 2 and row[2] == 800 and row[3] == 200_000.0
    assert section(report, "By client")["rows"][0][2] == 200_000.0
    assert "whole months" in report["notes"][0]


async def test_the_fuel_report_ranks_vehicles_by_kilometres_a_litre(client):
    f = await fleet(client)
    await history(f, [(1000, 250, 0, None), (500, 150, 0, None)])  # 1,500 km on 400 litres at KES 180
    report = await run(client, f.owner, "fuel_efficiency")
    [row] = section(report, "By vehicle")["rows"]
    assert row[:6] == ["KCA 123A", 2, 1500, 400.0, 3.75, 72_000.0] and row[6] == 48.0 and row[7] == 0.0  # KES 48 of fuel for every km
    assert section(report, "By driver")["rows"] == []  # these trips had no driver recorded


async def test_the_debtors_report_is_a_picture_of_now_with_ageing(client):
    owner, vehicle, _ = await lease_setup(client)
    c = await seed_client("Slow Payer Ltd")
    await seed_trip(vehicle["id"], revenue_cents=5_000_000, client_id=c, at=moment(-3, 2), km=300)
    report = await run(client, owner, "debtors")
    rows = section(report, "Clients")["rows"]
    assert rows[0][0] == "Slow Payer Ltd" and rows[0][1] == 50_000.0 and rows[0][-2] == 1 and rows[0][-1] > 30
    assert rows[-1][0] == "Everyone" and rows[-1][1] == 50_000.0 and "picture of now" in report["notes"][0]
    assert report["from"] == report["to"]


async def test_the_alert_history_lists_every_alert_and_how_often_each_check_fired(client):
    f = await seeded(client)
    await drive(client, f, litres=400, km=1000)
    report = await run(client, f.owner, "alerts", start=(nairobi_today() - timedelta(days=2)).isoformat(), end=nairobi_today().isoformat())
    kinds = section(report, "By kind")["rows"]
    assert kinds[0][:5] == ["Fuel above expected", 1, 1, 0, 0]
    [alert] = section(report, "Every alert")["rows"]
    assert alert[1] == "KCA 123A" and alert[2] == "driver user" and alert[3] == "red" and "more than expected" in alert[4] and alert[5] == "open"


async def test_the_scorecard_report_matches_the_scorecards(client):
    f = await seeded(client)
    await drive(client, f, litres=305, km=1000, phone_km=1000)
    today, yesterday = nairobi_today().isoformat(), (nairobi_today() - timedelta(days=1)).isoformat()  # the trip started 20 hours ago: yesterday, early in the day
    report = await run(client, f.owner, "scorecards", start=yesterday, end=today)
    page = (await client.get("/scorecards", params={"start": yesterday, "end": today}, headers=bearer(f.owner))).json()["drivers"][0]
    [row] = section(report, "Drivers")["rows"]
    assert row[0] == "driver user" and row[1] == page["trips"] and row[8] == page["overall"] and row[3] == page["safety"]


# ---- export ----------------------------------------------------------------------------------------------------------------------


async def test_a_report_is_exported_as_excel_with_a_sheet_per_section_and_as_pdf(client):
    f = await fleet(client)
    await history(f, [(1000, 250, 0, None)])
    xl = await client.get("/report-catalog/fuel_efficiency/export", params={**span(), "file_format": "xlsx"}, headers=bearer(f.owner))
    assert xl.status_code == 200 and xl.headers["content-type"].startswith("application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    assert 'filename="fleettms-fuel-efficiency-' in xl.headers["content-disposition"] and xl.headers["content-disposition"].endswith('.xlsx"')
    book = load_workbook(io.BytesIO(xl.content))
    assert book.sheetnames[0].startswith("By vehicle") and [c.value for c in book[book.sheetnames[0]][1]][:2] == ["Vehicle", "Trips"] and book[book.sheetnames[0]]["A2"].value == "KCA 123A"
    pdf = await client.get("/report-catalog/fuel_efficiency/export", params=span(), headers=bearer(f.owner))
    assert pdf.content.startswith(b"%PDF") and pdf.headers["content-type"] == "application/pdf" and pdf.headers["content-disposition"].endswith('.pdf"')
    assert (await client.get("/report-catalog/fuel_efficiency/export", params={**span(), "file_format": "docx"}, headers=bearer(f.owner))).status_code == 422
    assert (await client.get("/report-catalog/profit/export", headers=bearer((await staff_session(client, f.owner, "supervisor", "sup@example.com"))[0]))).status_code == 403


async def test_every_report_in_the_catalogue_runs_and_exports_on_an_empty_business(client):
    owner, _ = await owner_session(client)
    for key in ("summary", "profit", "fuel_efficiency", "debtors", "alerts", "scorecards"):
        assert (await client.get(f"/report-catalog/{key}", params=span(), headers=bearer(owner))).status_code == 200, key
        for fmt in ("pdf", "xlsx"):
            res = await client.get(f"/report-catalog/{key}/export", params={**span(), "file_format": fmt}, headers=bearer(owner))
            assert res.status_code == 200 and len(res.content) > 500, (key, fmt)


async def test_one_business_never_sees_anothers_figures_in_a_report(client):
    f = await fleet(client)
    await history(f, [(1000, 250, 0, None)])
    other, _ = await owner_session(client, "Bravo", "b@example.com")
    assert section(await run(client, other, "fuel_efficiency"), "By vehicle")["rows"] == []


# ---- plans -----------------------------------------------------------------------------------------------------------------------


async def test_starter_keeps_the_basic_summary_and_the_rest_come_with_standard(client, billing):
    from tests.billing_helpers import pay_invoice

    f = await fleet(client)
    await pay_invoice(client, f.owner)
    await client.put("/subscription/plans", headers=bearer(f.owner), json={"vehicles": [{"vehicle_id": f.vehicle["id"], "plan": "starter"}]})
    listed = {r["key"]: r for r in (await client.get("/report-catalog", headers=bearer(f.owner))).json()}
    assert listed["summary"]["available"] is True and listed["profit"]["available"] is False and listed["profit"]["plan_needed"] == "standard"
    assert (await client.get("/report-catalog/summary", params=span(), headers=bearer(f.owner))).status_code == 200
    gated = await client.get("/report-catalog/profit", params=span(), headers=bearer(f.owner))
    assert gated.status_code == 402 and gated.json()["detail"]["code"] == "plan_required"
    await client.put("/subscription/plans", headers=bearer(f.owner), json={"vehicles": [{"vehicle_id": f.vehicle["id"], "plan": "standard"}]})
    assert (await client.get("/report-catalog/profit", params=span(), headers=bearer(f.owner))).status_code == 200


# ---- scheduled ---------------------------------------------------------------------------------------------------------------------


async def test_a_schedule_can_choose_the_report_and_the_file_format(client):
    f = await fleet(client)
    await history(f, [(1000, 250, 0, None)])
    made = await client.post("/report-schedules", headers=bearer(f.owner), json={"frequency": "monthly", "channel": "email", "recipient": "boss@example.com", "report": "fuel_efficiency", "file_format": "xlsx"})
    assert made.status_code == 201 and made.json()["report"] == "fuel_efficiency" and made.json()["file_format"] == "xlsx"
    same_but_pdf = await client.post("/report-schedules", headers=bearer(f.owner), json={"frequency": "monthly", "channel": "email", "recipient": "boss@example.com", "report": "fuel_efficiency", "file_format": "pdf"})
    assert same_but_pdf.status_code == 201  # a different file is a different schedule
    dup = await client.post("/report-schedules", headers=bearer(f.owner), json={"frequency": "monthly", "channel": "email", "recipient": "boss@example.com", "report": "fuel_efficiency", "file_format": "xlsx"})
    assert dup.status_code == 409
    sent = await client.post(f"/report-schedules/{made.json()['id']}/send-now", headers=bearer(f.owner))
    assert sent.status_code == 200 and sent.json() == {"sent": True}
    [mail] = fake_sender("email").outbox
    assert mail["mime"] == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet" and mail["filename"].endswith(".xlsx") and "Fuel efficiency" in mail["subject"]
    assert load_workbook(io.BytesIO(mail["pdf"])).sheetnames[0].startswith("By vehicle")
    assert await send_scheduled_reports(date(2026, 11, 3)) >= 1  # the monthly run sends each schedule too
    assert {m["mime"] for m in fake_sender("email").outbox} == {"application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "application/pdf"}


async def test_a_schedule_checks_the_report_and_who_may_run_it(client):
    f = await fleet(client)
    body = {"frequency": "daily", "channel": "email", "recipient": "boss@example.com"}
    assert (await client.post("/report-schedules", headers=bearer(f.owner), json={**body, "report": "nonsense"})).json()["detail"]["code"] == "unknown_report"
    assert (await client.post("/report-schedules", headers=bearer(f.owner), json={**body, "file_format": "docx"})).status_code == 422
    manager, _ = await staff_session(client, f.owner, "manager", "mgr@example.com")
    assert (await client.post("/report-schedules", headers=bearer(manager), json={**body, "report": "profit"})).status_code == 403  # a manager does not see profit
    assert (await client.post("/report-schedules", headers=bearer(manager), json={**body, "report": "fuel_efficiency"})).status_code == 201
